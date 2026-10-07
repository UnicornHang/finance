"""发票归档导出：纯内存生成 `.xlsx`（无 DB / Celery）。

列顺序与设计 §5.1、前端归档导出字段对齐。
调用方负责把类型/状态译为中文标签、归档时间为本地化字符串；本模块只写表。
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font

# 表头（顺序即列顺序）
_HEADERS: tuple[str, ...] = (
    "发票抬头",
    "开票公司",
    "纳税人识别号",
    "发票代码",
    "发票号码",
    "开票日期",
    "不含税金额",
    "税额",
    "含税合计",
    "发票类型",
    "销售方",
    "购买方",
    "备注",
    "状态",
    "操作用户",
    "归档时间",
)

# 金额列在工作表中的 1-based 列号
_AMOUNT_COLS: frozenset[int] = frozenset({7, 8, 9})


@dataclass(frozen=True, slots=True)
class InvoiceRow:
    """已解析好的发票导出行（标签与时间由上游本地化）。"""

    invoice_title: str | None = None
    company: str | None = None
    tax_id: str | None = None
    invoice_code: str | None = None
    invoice_number: str | None = None
    invoice_date: str | None = None
    amount_excl_tax: Decimal | float | int | None = None
    tax_amount: Decimal | float | int | None = None
    amount_incl_tax: Decimal | float | int | None = None
    invoice_type_label: str | None = None
    seller: str | None = None
    buyer: str | None = None
    remark: str | None = None
    status_label: str | None = None
    operator_name: str | None = None
    created_at_label: str | None = None


def _cell_text(value: object | None) -> str:
    """空值转为空串，其余转字符串。"""
    if value is None:
        return ""
    return str(value)


def _cell_number(value: Decimal | float | int | None) -> float | None:
    """金额转 float；空值保持空单元格。"""
    if value is None:
        return None
    if isinstance(value, Decimal):
        return float(value)
    return float(value)


def build_invoice_xlsx(rows: Sequence[InvoiceRow]) -> bytes:
    """将发票行写成 xlsx 字节；允许空 rows（仅表头）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "发票归档"

    header_font = Font(bold=True)
    for col_idx, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font

    for row_idx, row in enumerate(rows, start=2):
        values: list[object | None] = [
            _cell_text(row.invoice_title),
            _cell_text(row.company),
            _cell_text(row.tax_id),
            _cell_text(row.invoice_code),
            _cell_text(row.invoice_number),
            _cell_text(row.invoice_date),
            _cell_number(row.amount_excl_tax),
            _cell_number(row.tax_amount),
            _cell_number(row.amount_incl_tax),
            _cell_text(row.invoice_type_label),
            _cell_text(row.seller),
            _cell_text(row.buyer),
            _cell_text(row.remark),
            _cell_text(row.status_label),
            _cell_text(row.operator_name),
            _cell_text(row.created_at_label),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            # 金额列显式为数值，避免被当成文本
            if col_idx in _AMOUNT_COLS and isinstance(value, (int, float)):
                cell.number_format = "0.00"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
