"""发票 xlsx / 合同 PDF+zip 生成器单测（纯内存，无 DB）。"""

from __future__ import annotations

import zipfile
from io import BytesIO

from openpyxl import load_workbook

from app.services.export_contract_pdf import (
    ContractView,
    ViolationItem,
    build_contract_report_pdf,
    build_contract_reports_zip,
    safe_pdf_filename,
)
from app.services.export_invoice_xlsx import InvoiceRow, build_invoice_xlsx


def test_invoice_xlsx_has_header_and_row():
    """表头与首行数据对齐设计 §5.1；金额为数值类型。"""
    data = build_invoice_xlsx(
        [
            InvoiceRow(
                invoice_title="抬头",
                company="公司",
                tax_id="91110000",
                invoice_code="044001",
                invoice_number="12345678",
                invoice_date="2026-10-01",
                amount_excl_tax=1000.0,
                tax_amount=130.0,
                amount_incl_tax=1130.0,
                invoice_type_label="电子发票",
                seller="销方",
                buyer="购方",
                remark="备注",
                status_label="已归档",
                operator_name="张三",
                created_at_label="2026/10/07 12:00",
            )
        ]
    )
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "发票抬头"
    assert ws["P1"].value == "归档时间"
    assert ws["A2"].value == "抬头"
    assert ws["J2"].value == "电子发票"
    assert ws["N2"].value == "已归档"
    assert ws["P2"].value == "2026/10/07 12:00"
    assert ws["I2"].value == 1130.0
    assert isinstance(ws["I2"].value, (int, float))


def test_invoice_xlsx_empty_rows_still_has_header():
    """空结果仍出仅含表头的 xlsx。"""
    data = build_invoice_xlsx([])
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "发票抬头"
    assert ws["A2"].value is None


def test_contract_pdf_without_review_still_builds():
    """无审查结果仍生成合法 PDF（%PDF 头）。"""
    pdf = build_contract_report_pdf(
        ContractView(
            contract_name="测试合同",
            contract_no="C-1",
            party_a="甲",
            party_b="乙",
            risk_level="high",
            status="active",
            summary=None,
            violations=[],
            exported_at="2026/10/07 12:00",
        )
    )
    assert pdf[:4] == b"%PDF"


def test_contract_pdf_with_violations_builds():
    """含摘要与违规项的 PDF 可生成。"""
    pdf = build_contract_report_pdf(
        ContractView(
            contract_name="有审查合同",
            contract_no="C-2",
            party_a="甲方公司",
            party_b="乙方公司",
            sign_date="2026-01-01",
            effective_start="2026-01-02",
            amount=100000,
            risk_level="medium",
            status="active",
            summary="存在付款条款风险",
            violations=[
                ViolationItem(clause="第3条", issue="付款周期过长", severity="high"),
                ViolationItem(clause="第5条", issue="违约金过高", severity="medium"),
            ],
            exported_at="2026/10/07 12:30",
        )
    )
    assert pdf[:4] == b"%PDF"
    assert len(pdf) > 200


def test_zip_contains_n_pdfs():
    """zip 内文件数等于 PDF 份数；N=1 仍为 zip。"""
    pdf_a = build_contract_report_pdf(
        ContractView(contract_name="A", contract_no="A-1", summary=None, violations=[])
    )
    pdf_b = build_contract_report_pdf(
        ContractView(contract_name="B", contract_no="B-1", summary="有摘要", violations=[])
    )
    z = build_contract_reports_zip(
        [
            ("a-11111111.pdf", pdf_a),
            ("b-22222222.pdf", pdf_b),
        ]
    )
    with zipfile.ZipFile(BytesIO(z)) as zf:
        names = zf.namelist()
        assert len(names) == 2
        assert "a-11111111.pdf" in names
        assert "b-22222222.pdf" in names
        assert zf.read("a-11111111.pdf")[:4] == b"%PDF"

    z1 = build_contract_reports_zip([("only-aaaaaaaa.pdf", pdf_a)])
    with zipfile.ZipFile(BytesIO(z1)) as zf:
        assert len(zf.namelist()) == 1


def test_empty_zip_is_valid():
    """空合同列表仍生成合法空 zip。"""
    z = build_contract_reports_zip([])
    with zipfile.ZipFile(BytesIO(z)) as zf:
        assert zf.namelist() == []


def test_safe_pdf_filename_strips_illegal_chars():
    """文件名非法字符替换，并附短 id。"""
    name = safe_pdf_filename(r'合/同:名*称?', "abcd1234")
    assert name == "合_同_名_称-abcd1234.pdf"
    assert all(c not in name for c in '\\/:*?"<>|')
    # 输入已带 .pdf 时不重复后缀
    assert safe_pdf_filename("报告.pdf", "deadbeef") == "报告-deadbeef.pdf"
